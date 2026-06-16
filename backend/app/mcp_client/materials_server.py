"""강의자료 RAG MCP 서버 (materials__*).

services/material_service.py 가 Obsidian Vault('강의자료/{과목}/…')에 받아둔 원본
강의자료(PPT/PDF/문서)의 **본문 텍스트**를 추출해 학습 비서가 내용 기반으로 답하게 한다.
(기존엔 비서가 자료 본문을 못 읽었다 — uiConfig 주석 참고.)

도구
----
- materials__list_files : 받아둔 강의자료 파일 목록(과목별)
- materials__read       : 한 파일의 본문 텍스트 추출(요약은 LLM 몫 — 원문만 제공)
- materials__search     : 키워드로 자료 본문 검색 → 매칭 파일 + 스니펫

⚠️ 역할 분리: 추출·검색만 한다(요약/RAG 추론은 챗 LLM). Obsidian 에 받아둔 파일이
소스라 setup.py 는 obsidian 설정 시에만 마운트한다(obsidian 과 동일 게이트).
텍스트 추출(pypdf/python-pptx/python-docx)과 Obsidian REST(_list_dir/_read_bytes)는
모듈 레벨 함수라 tests/test_materials_mcp.py 가 오프라인으로 mock 한다.
"""
import io
import json
from urllib.parse import quote, urlparse

import httpx
from mcp.server import Server
from mcp.types import TextContent, Tool

MATERIALS_ROOT = "강의자료"
MAX_TEXT = 12000          # read 가 돌려주는 본문 최대 길이(토큰 보호)
SEARCH_FILE_CAP = 50      # search 가 본문 추출까지 하는 파일 수 상한
SNIPPET_PAD = 80          # 검색 스니펫 앞뒤 문맥 길이
_LOCAL_HOSTS = {"localhost", "127.0.0.1", "::1"}


# ── 텍스트 추출 (순수, bytes → str) ──────────────────────────────────────────
def _extract_text(filename: str, data: bytes) -> str:
    """확장자별 본문 텍스트 추출. 미지원/실패 시 빈 문자열에 가깝게(안내) 반환."""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    try:
        if ext == "pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            return "\n".join((p.extract_text() or "") for p in reader.pages)
        if ext == "pptx":
            from pptx import Presentation
            prs = Presentation(io.BytesIO(data))
            parts = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if shape.has_text_frame and shape.text_frame.text.strip():
                        parts.append(shape.text_frame.text)
            return "\n".join(parts)
        if ext == "docx":
            from docx import Document
            doc = Document(io.BytesIO(data))
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        if ext == "ipynb":
            nb = json.loads(data.decode("utf-8", errors="ignore"))
            parts = []
            for cell in nb.get("cells", []):
                src = cell.get("source", [])
                parts.append("".join(src) if isinstance(src, list) else str(src))
            return "\n".join(parts)
        # 그 외(txt/md/py 등)는 평문 시도
        return data.decode("utf-8", errors="ignore")
    except Exception as e:  # 손상 파일 등 — 도구가 죽지 않게 안내 텍스트로
        return f"(본문 추출 실패: {e})"


def _snippets(text: str, query: str, max_hits: int = 3) -> list[str]:
    """본문에서 query 주변 스니펫 추출(대소문자 무시). 없으면 빈 리스트."""
    low = text.lower()
    q = query.lower()
    out, start = [], 0
    while len(out) < max_hits:
        i = low.find(q, start)
        if i < 0:
            break
        a = max(0, i - SNIPPET_PAD)
        b = min(len(text), i + len(query) + SNIPPET_PAD)
        out.append(" ".join(text[a:b].split()))
        start = b
    return out


def _tls_verify(base_url: str) -> bool:
    return urlparse(base_url).hostname not in _LOCAL_HOSTS


# ── Obsidian REST IO (테스트 monkeypatch 경계) ───────────────────────────────
async def _list_dir(base: str, headers: dict, verify: bool, rel_dir: str) -> list[str]:
    """Vault 의 디렉터리 항목 목록(파일은 이름, 폴더는 '이름/'). 실패 시 빈 리스트."""
    url = f"{base}/vault/{quote(rel_dir.strip('/'))}/"
    try:
        async with httpx.AsyncClient(verify=verify, timeout=10) as c:
            r = await c.get(url, headers=headers)
            if r.status_code == 200:
                return r.json().get("files", [])
    except Exception:
        pass
    return []


async def _read_bytes(base: str, headers: dict, verify: bool, rel_path: str) -> bytes | None:
    """Vault 파일 원본 바이트. 실패 시 None."""
    url = f"{base}/vault/{quote(rel_path.strip('/'))}"
    try:
        async with httpx.AsyncClient(verify=verify, timeout=20) as c:
            r = await c.get(url, headers=headers)
            if r.status_code == 200:
                return r.content
    except Exception:
        pass
    return None


# ── 파일 수집 / 디스패치 ─────────────────────────────────────────────────────
def _root(vault_path: str) -> str:
    vp = (vault_path or "").strip("/")
    return f"{vp}/{MATERIALS_ROOT}" if vp else MATERIALS_ROOT


async def _collect_files(cfg: dict, course: str | None = None) -> list[dict]:
    """받아둔 강의자료 파일 목록 [{course, file, path}]. course 지정 시 그 과목만."""
    base, headers, verify = cfg["base"], cfg["headers"], cfg["verify"]
    root = _root(cfg["vault"])
    if course:
        courses = [course if course.endswith("/") else course + "/"]
    else:
        courses = [f for f in await _list_dir(base, headers, verify, root) if f.endswith("/")]
    out = []
    for cdir in courses:
        cname = cdir.rstrip("/")
        files = await _list_dir(base, headers, verify, f"{root}/{cname}")
        for f in files:
            if not f.endswith("/"):
                out.append({"course": cname, "file": f, "path": f"{MATERIALS_ROOT}/{cname}/{f}"})
    return out


def _full_path(cfg: dict, rel_path: str) -> str:
    """'강의자료/…' 상대경로 → vault_path 포함 전체 경로."""
    vp = (cfg["vault"] or "").strip("/")
    rel = rel_path.strip("/")
    return f"{vp}/{rel}" if vp else rel


async def _dispatch(name: str, arguments: dict, cfg: dict) -> list[TextContent]:
    arguments = arguments or {}

    if name == "list_files":
        files = await _collect_files(cfg, arguments.get("course"))
        return [TextContent(type="text", text=json.dumps(files, ensure_ascii=False))]

    if name == "read":
        rel = arguments["path"]
        data = await _read_bytes(cfg["base"], cfg["headers"], cfg["verify"], _full_path(cfg, rel))
        if data is None:
            return [TextContent(type="text", text=f"(파일을 찾을 수 없습니다: {rel})")]
        text = _extract_text(rel.rsplit("/", 1)[-1], data)
        if len(text) > MAX_TEXT:
            text = text[:MAX_TEXT] + f"\n…(이하 생략 — 총 {len(text)}자)"
        return [TextContent(type="text", text=text or "(추출된 텍스트가 없습니다)")]

    if name == "search":
        query = (arguments.get("query") or "").strip()
        if not query:
            raise ValueError("search 에는 query 가 필요합니다.")
        files = await _collect_files(cfg, arguments.get("course"))
        capped = files[:SEARCH_FILE_CAP]
        hits = []
        for f in capped:
            data = await _read_bytes(cfg["base"], cfg["headers"], cfg["verify"],
                                     _full_path(cfg, f["path"]))
            if data is None:
                continue
            snips = _snippets(_extract_text(f["file"], data), query)
            if snips:
                hits.append({"path": f["path"], "course": f["course"], "snippets": snips})
        result = {"query": query, "matched": len(hits), "scanned": len(capped),
                  "results": hits}
        if len(files) > SEARCH_FILE_CAP:
            result["note"] = f"파일 {len(files)}개 중 {SEARCH_FILE_CAP}개만 검색(상한)."
        return [TextContent(type="text", text=json.dumps(result, ensure_ascii=False))]

    raise ValueError(f"unknown tool: {name}")


def create_materials_mcp_server(auth_code: str, base_url: str, vault_path: str = "") -> Server:
    """강의자료 RAG MCP 서버 생성.

    auth_code/base_url: Obsidian Local REST API (받아둔 파일 읽기용).
    vault_path: vault 내부 상대 경로(보통 'LMS-Bridge' 또는 빈 문자열).
    """
    server = Server("materials-mcp")
    base = base_url.rstrip("/")
    cfg = {
        "base": base,
        "headers": {"Authorization": f"Bearer {auth_code}"},
        "verify": _tls_verify(base_url),
        "vault": vault_path,
    }

    @server.list_tools()
    async def list_tools():
        return [
            Tool(
                name="list_files",
                description=(
                    "Obsidian 에 받아둔 강의자료 파일 목록을 반환(과목별). "
                    "course(과목 폴더명) 지정 시 해당 과목만. 반환: [{course,file,path}] JSON. "
                    "read/search 의 path 를 여기서 얻는다."
                ),
                inputSchema={"type": "object", "properties": {"course": {"type": "string"}}},
            ),
            Tool(
                name="read",
                description=(
                    "강의자료 파일 1개의 본문 텍스트를 추출해 반환(PPT/PDF/문서). "
                    "path 는 list_files 의 path('강의자료/{과목}/{파일}'). "
                    "요약·정리는 이 본문을 바탕으로 네가(LLM) 해라."
                ),
                inputSchema={
                    "type": "object",
                    "properties": {"path": {"type": "string"}},
                    "required": ["path"],
                },
            ),
            Tool(
                name="search",
                description=(
                    "강의자료 본문에서 키워드를 검색해 매칭 파일과 스니펫을 반환. "
                    "'신경망 어느 자료에 나와?', '시험 범위 관련 슬라이드 찾아줘' 같은 질문에 사용. "
                    "course 로 과목을 좁힐 수 있다. 반환: {matched, results:[{path,snippets}]} JSON."
                ),
                inputSchema={
                    "type": "object",
                    "properties": {
                        "query": {"type": "string"},
                        "course": {"type": "string"},
                    },
                    "required": ["query"],
                },
            ),
        ]

    @server.call_tool()
    async def call_tool(name: str, arguments: dict):
        return await _dispatch(name, arguments, cfg)

    return server
